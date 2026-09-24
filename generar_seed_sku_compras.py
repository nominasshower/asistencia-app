"""Genera sku_compras_seed.sql: catálogo Familia/Subfamilia + SKUs existentes.
Fuente de verdad por campo:
  - Odoo (vía XML-RPC, solo lectura): qué SKUs existen HOY y en qué empresa.
  - Excel "Plantilla Suministro": precio estándar, moneda, tipo de producto y si es
    sustancia peligrosa (el pull de Odoo no trae esos campos). NO se usan las columnas
    de código por planta (ARTICULO VIA MVT/SW/MW/KSK) — esas ya no se usan, todo vive
    unificado en Odoo.
Re-ejecutable: el SQL usa ON CONFLICT, sirve para re-sincronizar altas hechas directo en Odoo."""
import collections
import re
import xmlrpc.client

import openpyxl

EXCEL = r'D:\mis documentos\Odoo - Listado de Articulos - Relacion Articulo VIA MVT-SW-MW-KSK 260721.xlsx'
OUT = r'D:\asistencia-app\sku_compras_seed.sql'
ODOO_URL = 'https://swmvt24-swmvt.odoo.com'
ODOO_DB = 'swmvt24-swmvt-swmvt-12590405'
ODOO_USER = 'sistemas@modular-tops.com'
ODOO_PASS = 'Modular-tops.s1'

PAT = re.compile(r'^([A-Z]{2,3}-\d{2})-(\d{3,5})$')
CODIGO = re.compile(r'^[A-Z]{2,3}-\d{2}$')  # descarta typos del Excel como 'EM-88-'
TIPOS = {'ALMACENABLE', 'CONSUMIBLE', 'SERVICIO'}


def q(v):
    return 'NULL' if v is None else "'" + str(v).replace("'", "''") + "'"


def qn(v):
    return 'NULL' if v is None else str(v)


def leer_catalogo():
    ws = openpyxl.load_workbook(EXCEL, data_only=True)['Familia-Subfamilia']
    fam = sub = None
    out = []
    for r in range(5, ws.max_row + 1):
        f, s, c = (ws.cell(row=r, column=i).value for i in (1, 2, 3))
        if f == '(en blanco)' or f == 'Total general':
            break
        fam = (f or fam).strip()
        sub = (s or sub).strip()
        c = (c or '').strip()
        if CODIGO.match(c):
            out.append((c, fam, sub, len(out)))
    return out


def leer_plantilla_suministro():
    """sku -> (tipo_producto, precio, moneda, sustancia_peligrosa)"""
    ws = openpyxl.load_workbook(EXCEL, data_only=True)['Plantilla Suministro']
    out = {}
    for r in range(7, ws.max_row + 1):
        sku = ws.cell(row=r, column=2).value
        if not isinstance(sku, str):
            continue
        sku = sku.strip().upper()
        if not PAT.match(sku):
            continue
        tipo = (ws.cell(row=r, column=5).value or '').strip().upper() or None
        if tipo not in TIPOS:
            tipo = None
        precio = ws.cell(row=r, column=8).value
        try:
            precio = float(str(precio).replace(',', '')) if precio not in (None, '') else None
        except ValueError:
            precio = None
        moneda = (ws.cell(row=r, column=14).value or '').strip().upper() or None
        peligrosa = (ws.cell(row=r, column=15).value or '').strip().upper() == 'SUSTANCIA PELIGROSA'
        out[sku] = (tipo, precio, moneda, peligrosa)
    return out


def leer_odoo():
    uid = xmlrpc.client.ServerProxy(ODOO_URL + '/xmlrpc/2/common').authenticate(ODOO_DB, ODOO_USER, ODOO_PASS, {})
    m = xmlrpc.client.ServerProxy(ODOO_URL + '/xmlrpc/2/object', allow_none=True)
    return m.execute_kw(ODOO_DB, uid, ODOO_PASS, 'product.product', 'search_read',
                        [[['default_code', '!=', False], '|', ['active', '=', True], ['active', '=', False]]],
                        {'fields': ['default_code', 'name', 'company_id', 'uom_id', 'categ_id'],
                         'context': {'active_test': False}})


def main():
    catalogo = leer_catalogo()
    codigos = {c for c, *_ in catalogo}
    extra = leer_plantilla_suministro()
    productos = []
    categorias = collections.defaultdict(collections.Counter)
    for p in leer_odoo():
        sku = p['default_code'].strip().upper()
        m = PAT.match(sku)
        if not m or m.group(1) not in codigos:
            continue
        productos.append((sku, m.group(1), int(m.group(2)), p, extra.get(sku)))
        if p['categ_id']:
            categorias[m.group(1)][p['categ_id'][1]] += 1

    con_extra = sum(1 for *_, e in productos if e)
    lines = ['-- Generado por generar_seed_sku_compras.py — correr DESPUÉS de sku_compras_setup.sql', 'BEGIN;', '']
    lines.append('INSERT INTO sku_catalogo (codigo, familia, subfamilia, categoria_odoo, orden) VALUES')
    vals = []
    for c, f, s, orden in catalogo:
        cat = categorias[c].most_common(1)[0][0] if categorias[c] else None
        vals.append(f'  ({q(c)}, {q(f)}, {q(s)}, {q(cat)}, {orden})')
    lines.append(',\n'.join(vals))
    lines.append('ON CONFLICT (codigo) DO UPDATE SET familia=EXCLUDED.familia, subfamilia=EXCLUDED.subfamilia,')
    lines.append('  categoria_odoo=COALESCE(EXCLUDED.categoria_odoo, sku_catalogo.categoria_odoo), orden=EXCLUDED.orden;')
    lines.append('')

    for i in range(0, len(productos), 400):
        lines.append('INSERT INTO sku_articulos (sku, codigo, consecutivo, nombre, um, empresa, '
                      'tipo_producto, precio_estandar, moneda, sustancia_peligrosa, origen, en_odoo) VALUES')
        vals = []
        for sku, cod, n, p, ex in productos[i:i + 400]:
            emp = p['company_id'][1] if p['company_id'] else 'COMPARTIDO'
            um = p['uom_id'][1] if p['uom_id'] else None
            tipo, precio, moneda, peligrosa = ex or (None, None, None, False)
            vals.append(f"  ({q(sku)}, {q(cod)}, {n}, {q(p['name'].strip())}, {q(um)}, {q(emp)}, "
                        f"{q(tipo or 'ALMACENABLE')}, {qn(precio)}, {q(moneda)}, {str(peligrosa).lower()}, 'ODOO', true)")
        lines.append(',\n'.join(vals))
        lines.append('ON CONFLICT (sku) DO UPDATE SET tipo_producto=EXCLUDED.tipo_producto, '
                      'precio_estandar=EXCLUDED.precio_estandar, moneda=EXCLUDED.moneda, '
                      'sustancia_peligrosa=EXCLUDED.sustancia_peligrosa, empresa=EXCLUDED.empresa;')
        lines.append('')
    lines.append('COMMIT;')
    open(OUT, 'w', encoding='utf-8').write('\n'.join(lines))
    print(f'catalogo={len(catalogo)} codigos, skus_odoo={len(productos)}, con_datos_excel={con_extra} -> {OUT}')


if __name__ == '__main__':
    main()
